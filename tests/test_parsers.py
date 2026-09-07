import pytest
from pathlib import Path
from parsers import parse_pdf
from extractor import extract_email, extract_github, extract_candidate_name

def test_malformed_pdf_does_not_crash(tmp_path):
    corrupt_file = tmp_path / "corrupted.pdf"
    corrupt_file.write_bytes(b"NOT A REAL PDF CONTENT RANDOM BYTES")
    
    candidate = parse_pdf(corrupt_file)
    assert candidate.parse_error is not None
    assert "PDF parsing error" in candidate.parse_error
    assert candidate.raw_text == ""

def test_empty_pdf_handled(tmp_path):
    empty_file = tmp_path / "empty.pdf"
    empty_file.write_bytes(b"")

    candidate = parse_pdf(empty_file)
    assert candidate.parse_error is not None
    assert "empty" in candidate.parse_error.lower()

def test_contact_info_extraction():
    sample_text = """
    Asha Rao
    asha.rao.dev@gmail.com | +91-9876543210 | github.com/asharao-tech
    Software Engineer with expertise in Python and LangGraph.
    """
    email = extract_email(sample_text)
    assert email == "asha.rao.dev@gmail.com"

    url, username = extract_github(sample_text)
    assert url == "https://github.com/asharao-tech"
    assert username == "asharao-tech"

    name = extract_candidate_name(sample_text, "candidate_asha.pdf")
    assert name == "Asha Rao"

def test_duplicate_resume_detection(tmp_path):
    """Verifies that duplicate resumes with different filenames are detected and only scored once."""
    from generate_sample_resumes import create_pdf
    from pipeline import process_resume_batch

    lines = [
        "Asha Rao",
        "asha@example.com",
        "SKILLS: Python, FastAPI, LangGraph, PostgreSQL",
        "PROJECTS: Autonomous agent pipeline using LangGraph and Python FastAPI."
    ]
    
    file1 = tmp_path / "resume_1.pdf"
    file2 = tmp_path / "resume_1_copy.pdf"
    create_pdf(file1, lines)
    create_pdf(file2, lines)

    results = process_resume_batch(tmp_path)
    assert results.batch_summary.total_resumes == 2
    assert results.batch_summary.duplicates == 1
    assert results.batch_summary.eligible == 1
    assert len(results.ranked_candidates) == 1
    assert results.ranked_candidates[0].candidate_name == "Asha Rao"
    assert any("Duplicate resume detected" in r for cand in results.rejected_candidates for r in cand.rejection_reasons)

def test_valid_docx_parsing(tmp_path):
    """Verifies that valid DOCX with paragraphs and table cells parses correctly."""
    import docx
    from parsers import parse_docx

    doc_path = tmp_path / "resume.docx"
    doc = docx.Document()
    doc.add_heading("Kiran Patel", level=0)
    doc.add_paragraph("kiran@example.com | github.com/kiran-ai")
    doc.add_paragraph("AI Engineer with expertise in Python and LangGraph.")

    # Add a table to verify table extraction
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skill Category"
    table.cell(0, 1).text = "Technologies"
    table.cell(1, 0).text = "Backend"
    table.cell(1, 1).text = "Python, FastAPI, Redis"
    doc.save(str(doc_path))

    candidate = parse_docx(doc_path)
    assert candidate.parse_error is None
    assert "Kiran Patel" in candidate.raw_text
    assert "kiran@example.com" in candidate.raw_text
    assert "Python, FastAPI, Redis" in candidate.raw_text

def test_corrupt_docx_handled(tmp_path):
    """Verifies that a malformed/corrupt DOCX file returns a parse error without crashing."""
    from parsers import parse_docx

    corrupt_file = tmp_path / "corrupted.docx"
    corrupt_file.write_bytes(b"INVALID_NON_ZIP_DOCX_HEADER_DATA")

    candidate = parse_docx(corrupt_file)
    assert candidate.parse_error is not None
    assert "DOCX parsing error" in candidate.parse_error
    assert candidate.raw_text == ""

def test_empty_docx_handled(tmp_path):
    """Verifies that a 0-byte DOCX file returns an empty file error."""
    from parsers import parse_docx

    empty_file = tmp_path / "empty.docx"
    empty_file.write_bytes(b"")

    candidate = parse_docx(empty_file)
    assert candidate.parse_error is not None
    assert "empty" in candidate.parse_error.lower()

def test_docx_discovered_and_processed_by_pipeline(tmp_path):
    """Verifies that .docx files are discovered by the pipeline alongside error handling."""
    from generate_sample_resumes import create_docx
    from pipeline import process_resume_batch

    valid_docx = tmp_path / "candidate_kiran.docx"
    create_docx(
        valid_docx,
        lines=[
            "Kiran Patel",
            "kiran@example.com",
            "SKILLS: Python, FastAPI, LlamaIndex, RAG, PostgreSQL",
            "PROJECTS: Built an enterprise RAG retrieval pipeline using Python and LlamaIndex."
        ]
    )

    # Add an empty docx to verify mixed batch resilience
    bad_docx = tmp_path / "candidate_bad.docx"
    bad_docx.write_bytes(b"")

    results = process_resume_batch(tmp_path)
    assert results.batch_summary.total_resumes == 2
    assert results.batch_summary.successfully_parsed == 1
    assert results.batch_summary.eligible == 1
    assert results.batch_summary.failed_or_unreadable == 1
    assert len(results.ranked_candidates) == 1
    assert results.ranked_candidates[0].candidate_name == "Kiran Patel"
