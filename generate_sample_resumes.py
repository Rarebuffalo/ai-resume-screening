from pathlib import Path

def create_pdf(file_path: Path, lines: list[str]) -> None:
    """Generates a minimal valid PDF document with text lines."""
    commands = ["BT", "/F1 11 Tf"]
    y = 750
    for line in lines:
        sanitized = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        commands.append(f"50 {y} Td ({sanitized}) Tj -50 -{y} Td")
        y -= 16
        if y < 50:
            break
    commands.append("ET")
    stream = "\n".join(commands)
    stream_bytes = stream.encode("latin1")
    stream_len = len(stream_bytes)

    obj1 = "1 0 obj <</Type /Catalog /Pages 2 0 R>> endobj\n"
    obj2 = "2 0 obj <</Type /Pages /Kids [3 0 R] /Count 1>> endobj\n"
    obj3 = "3 0 obj <</Type /Page /Parent 2 0 R /Resources <</Font <</F1 4 0 R>>>> /MediaBox [0 0 612 792] /Contents 5 0 R>> endobj\n"
    obj4 = "4 0 obj <</Type /Font /Subtype /Type1 /BaseFont /Helvetica>> endobj\n"
    obj5 = f"5 0 obj <</Length {stream_len}>> stream\n{stream}\nendstream endobj\n"

    header = "%PDF-1.4\n"
    pos1 = len(header)
    pos2 = pos1 + len(obj1)
    pos3 = pos2 + len(obj2)
    pos4 = pos3 + len(obj3)
    pos5 = pos4 + len(obj4)
    xref_pos = pos5 + len(obj5)

    xref = f"""xref
0 6
0000000000 65535 f 
{pos1:010d} 00000 n 
{pos2:010d} 00000 n 
{pos3:010d} 00000 n 
{pos4:010d} 00000 n 
{pos5:010d} 00000 n 
trailer <</Size 6 /Root 1 0 R>>
startxref
{xref_pos}
%%EOF"""

    file_path.parent.mkdir(parents=True, exist_ok=True)
    with open(file_path, "wb") as f:
        f.write((header + obj1 + obj2 + obj3 + obj4 + obj5 + xref).encode("latin1"))

def generate_all_samples(target_dir: Path, include_docx: bool = True):
    target_dir.mkdir(parents=True, exist_ok=True)

    # 1. Asha Rao - Top Candidate: Strong Python, FastAPI, LangGraph, Tools, GCP, Redis, Testing
    create_pdf(target_dir / "candidate_01_asha_rao.pdf", [
        "Asha Rao",
        "asha.rao@example.com | github.com/asharao | Bangalore, India",
        "",
        "SUMMARY",
        "Software engineering intern specializing in Python backend architecture and stateful AI agents.",
        "",
        "TECHNICAL SKILLS",
        "Languages: Python, SQL",
        "Backend & Frameworks: FastAPI, AsyncIO, PostgreSQL, Redis, SQLAlchemy, Celery",
        "AI & Agentic Systems: LangGraph, LangChain, Multi-Agent, Tool Calling, Vector Search, ChromaDB",
        "Cloud & DevOps: Docker, GCP, CI/CD, Pytest",
        "",
        "PROJECTS",
        "Autonomous Customer Support Agent (Python, FastAPI, LangGraph, Redis, GCP)",
        "- Architected a stateful multi-agent system using LangGraph and tool calling for incident routing.",
        "- Integrated ChromaDB vector store with hybrid retrieval and evaluation pipelines.",
        "- Engineered async FastAPI backend with PostgreSQL persistence and Redis distributed caching.",
        "- Deployed containerized services on GCP Cloud Run with automated Pytest CI/CD workflows."
    ])

    # 2. Rahul Sharma - Solid Candidate: Python, FastAPI, RAG, LlamaIndex, ChromaDB, Docker
    create_pdf(target_dir / "candidate_02_rahul_sharma.pdf", [
        "Rahul Sharma",
        "rahul.sharma@example.com | github.com/rahul-ai",
        "",
        "EDUCATION",
        "B.Tech in Computer Science",
        "",
        "SKILLS",
        "Python, FastAPI, LlamaIndex, RAG, ChromaDB, PostgreSQL, Docker, Pytest",
        "",
        "EXPERIENCE & PROJECTS",
        "Enterprise Document Search RAG Pipeline",
        "- Developed a Retrieval-Augmented Generation (RAG) system with LlamaIndex and ChromaDB.",
        "- Built REST API using Python FastAPI with PostgreSQL metadata storage.",
        "- Containerized application with Docker for seamless local deployment.",
        "- Implemented unit tests with Pytest covering embedding generation and retrieval queries."
    ])

    # 3. Dev Patel - Python Only (No AI/Agentic): Must be Rejected
    create_pdf(target_dir / "candidate_03_dev_patel.pdf", [
        "Dev Patel",
        "dev.patel@example.com | github.com/dev-backend",
        "",
        "PROFESSIONAL PROFILE",
        "Backend developer focusing on relational databases and asynchronous job queues.",
        "",
        "SKILLS",
        "Python, Django, PostgreSQL, Redis, Celery, Docker, Linux",
        "",
        "PROJECTS",
        "Scalable E-Commerce Backend Service",
        "- Built modular Django REST APIs with PostgreSQL database indexing and optimization.",
        "- Integrated Celery workers and Redis queue for background email notifications.",
        "- Containerized application stack using Docker Compose."
    ])

    # 4. Priya Nair - AI + Java/React (No Python): Must be Rejected
    create_pdf(target_dir / "candidate_04_priya_nair.pdf", [
        "Priya Nair",
        "priya.nair@example.com | github.com/priya-nair",
        "",
        "SKILLS",
        "Java, Spring Boot, React, Next.js, TypeScript, LangChain, Node.js",
        "",
        "PROJECTS",
        "Interactive React Support Chatbot",
        "- Developed responsive frontend dashboard using React and Next.js.",
        "- Implemented Node.js service using LangChain for customer FAQ interactions.",
        "- Connected with Java Spring Boot enterprise microservices."
    ])

    # 5. Kavita Verma - Mixed Stack (Java + React + Python + LangChain): Must be Accepted!
    create_pdf(target_dir / "candidate_05_kavita_verma.pdf", [
        "Kavita Verma",
        "kavita.verma@example.com | github.com/kavita-dev",
        "",
        "TECHNICAL EXPERTISE",
        "Languages: Python, Java, TypeScript",
        "Frameworks: FastAPI, Spring Boot, React, Next.js, LangChain",
        "Databases & Cloud: PostgreSQL, Docker, GCP, RAG",
        "",
        "PROJECTS",
        "Full-Stack Knowledge Engine (Python, FastAPI, LangChain, React, Docker)",
        "- Implemented document retrieval RAG service using Python and FastAPI.",
        "- Built interactive Next.js web application for knowledge search.",
        "- Migrated legacy Java Spring Boot services to Python microservices with Docker on GCP."
    ])

    # 6. Arjun Gupta - Thin LLM Wrapper (Python + OpenAI API call): Penalized
    create_pdf(target_dir / "candidate_06_arjun_gupta.pdf", [
        "Arjun Gupta",
        "arjun.gupta@example.com",
        "",
        "SKILLS",
        "Python, Flask, OpenAI API, HTML, CSS",
        "",
        "PROJECTS",
        "GPT-4 Question Answering App",
        "- Created simple wrapper around OpenAI API to generate responses to user questions.",
        "- Built simple Flask route that just calls the API with user prompt and displays output.",
        "- Hosted on Render free tier."
    ])

    # 7. Sam Wilson - Generic Tutorial Clone: Penalized
    create_pdf(target_dir / "candidate_07_sam_wilson.pdf", [
        "Sam Wilson",
        "sam.wilson@example.com",
        "",
        "SKILLS",
        "Python, LangChain, OpenAI API",
        "",
        "PROJECTS",
        "YouTube Summarizer Tutorial Clone",
        "- Completed youtube summarizer tutorial project from online bootcamp.",
        "- Implemented Titanic survival prediction baseline script."
    ])

    # 8. Corrupt PDF: Must be handled gracefully without crashing the batch
    corrupt_path = target_dir / "candidate_08_corrupt.pdf"
    with open(corrupt_path, "wb") as f:
        f.write(b"%PDF-1.4\nBROKEN_CORRUPT_DATA_NOT_A_VALID_OBJECT_STREAM\n%%EOF")

    # 9. Kiran Patel - Representative DOCX Candidate (Python, FastAPI, LlamaIndex, RAG, AWS)
    if include_docx:
        create_docx(
            target_dir / "candidate_09_kiran_patel.docx",
            lines=[
                "Kiran Patel",
                "kiran.patel@example.com | github.com/kiran-ai | Pune, India",
                "## PROFESSIONAL SUMMARY",
                "AI Systems Engineer with expertise in Python backend architectures and RAG pipelines.",
                "## TECHNICAL SKILLS",
                "Languages: Python, SQL",
                "Backend Frameworks: FastAPI, AsyncIO, PostgreSQL, Redis",
                "AI & LLM Stack: LlamaIndex, Vector Search, ChromaDB, RAG, Tool Calling",
                "Cloud & DevOps: Docker, AWS, Pytest, CI/CD",
                "## RECENT PROJECTS",
                "Autonomous Document RAG Engine (Python, FastAPI, LlamaIndex, ChromaDB)",
                "- Architected a Retrieval-Augmented Generation (RAG) system with ChromaDB vector search and LlamaIndex orchestration.",
                "- Engineered high-throughput async FastAPI REST API with PostgreSQL persistence and Redis caching.",
                "- Containerized microservices using Docker and deployed on AWS with automated Pytest suites."
            ],
            table_data=[
                ["Core Domain", "Technologies Used"],
                ["AI / Retrieval", "LlamaIndex, RAG, ChromaDB, Vector Search"],
                ["Backend / Cloud", "Python, FastAPI, PostgreSQL, Redis, Docker, AWS"]
            ]
        )
        print(f"Generated sample test resumes (8 PDFs + 1 DOCX) in '{target_dir}'.")
    else:
        print(f"Generated 8 sample test resumes in '{target_dir}'.")

def create_docx(file_path: Path, lines: list[str], table_data: list[list[str]] | None = None) -> None:
    """Generates a representative DOCX resume document."""
    import docx
    doc = docx.Document()
    for line in lines:
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.strip() == "":
            continue
        else:
            doc.add_paragraph(line)

    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, cell in enumerate(row):
                table.cell(r_idx, c_idx).text = cell

    file_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(file_path))

if __name__ == "__main__":
    generate_all_samples(Path("./resumes"))
