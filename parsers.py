import re
import logging
from pathlib import Path
from typing import Optional
from pypdf import PdfReader
from models import ExtractedCandidate

logger = logging.getLogger(__name__)

def clean_extracted_text(text: str) -> str:
    """Normalize whitespace and line breaks from extracted PDF text."""
    if not text:
        return ""
    # Replace multiple spaces with single space
    text = re.sub(r"[ \t]+", " ", text)
    # Normalize multiple newlines to at most two
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    # Strip leading/trailing whitespace
    return text.strip()

def parse_pdf(file_path: Path) -> ExtractedCandidate:
    """
    Parse a single PDF resume resiliently.
    Guarantees that an invalid or malformed PDF will never crash the caller.
    """
    file_name = file_path.name
    try:
        if not file_path.exists():
            return ExtractedCandidate(file_name=file_name, parse_error="File does not exist")

        if file_path.stat().st_size == 0:
            return ExtractedCandidate(file_name=file_name, parse_error="File is empty (0 bytes)")

        reader = PdfReader(str(file_path))
        if len(reader.pages) == 0:
            return ExtractedCandidate(file_name=file_name, parse_error="PDF contains 0 pages")

        page_texts = []
        for i, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text() or ""
                page_texts.append(page_text)
            except Exception as page_err:
                logger.warning(f"Error reading page {i} of {file_name}: {page_err}")
                continue

        full_text = "\n".join(page_texts)
        cleaned_text = clean_extracted_text(full_text)

        if not cleaned_text:
            return ExtractedCandidate(
                file_name=file_name,
                parse_error="No extractable text found (scanned or image-only PDF)"
            )

        return ExtractedCandidate(
            file_name=file_name,
            raw_text=cleaned_text
        )

    except Exception as e:
        logger.error(f"Failed to parse PDF {file_name}: {e}")
        return ExtractedCandidate(
            file_name=file_name,
            parse_error=f"PDF parsing error: {type(e).__name__}: {str(e)}"
        )

def parse_docx(file_path: Path) -> ExtractedCandidate:
    """
    Parse a single DOCX resume resiliently.
    Guarantees that an invalid, empty, or malformed DOCX will never crash the caller.
    Extracts text from both body paragraphs and table cells.
    """
    file_name = file_path.name
    try:
        if not file_path.exists():
            return ExtractedCandidate(file_name=file_name, parse_error="File does not exist")

        if file_path.stat().st_size == 0:
            return ExtractedCandidate(file_name=file_name, parse_error="File is empty (0 bytes)")

        import docx
        doc = docx.Document(str(file_path))

        chunks = []
        # 1. Extract paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                chunks.append(text)

        # 2. Extract table cells
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    deduped_cells = []
                    for c in row_cells:
                        if not deduped_cells or deduped_cells[-1] != c:
                            deduped_cells.append(c)
                    if deduped_cells:
                        chunks.append(" | ".join(deduped_cells))

        full_text = "\n".join(chunks)
        cleaned_text = clean_extracted_text(full_text)

        if not cleaned_text:
            return ExtractedCandidate(
                file_name=file_name,
                parse_error="No extractable text found in DOCX document"
            )

        return ExtractedCandidate(
            file_name=file_name,
            raw_text=cleaned_text
        )

    except Exception as e:
        logger.error(f"Failed to parse DOCX {file_name}: {e}")
        return ExtractedCandidate(
            file_name=file_name,
            parse_error=f"DOCX parsing error: {type(e).__name__}: {str(e)}"
        )
